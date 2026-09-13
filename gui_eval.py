"""
gui_eval.py
-----------
通知摘要评测 GUI 主程序（多模型并发版 + Prompt 编辑器）。

功能：
  1. [评测] 选择输入文件，配置多个评测模型并发评测，日志按模型分流显示
  2. [Prompt 编辑] 实时查看、编辑、保存当前 prompt 文件，自动同步字数统计

依赖：
    pip install openai openpyxl
"""

from __future__ import annotations

import json
import os
import queue
import re
import sys
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed, wait as cf_wait, FIRST_COMPLETED
from datetime import datetime
from pathlib import Path
from tkinter import (END, BooleanVar, IntVar, StringVar,
                     filedialog, messagebox, scrolledtext, ttk)
import tkinter as tk

ROOT = Path(__file__).parent
sys.path.insert(0, str(ROOT))

from core.excel_processor import (DataRow, ExcelOutputBuilder, ExcelReader,
                                   ExcelWriter, ModelEntry, ReviewerScore,
                                   detect_data_sheet, detect_sheets,
                                   _is_compact_format)

# ── 常量 ──────────────────────────────────────────────────────────────
PROMPT_FILE_DEFAULT = str(ROOT / "rules" / "llm_scoring_prompt.md")
OUTPUT_DIR_DEFAULT  = str(ROOT / "output")
VALID_EXTENSIONS    = {".xlsx", ".xls", ".csv"}
VALID_ISSUE_TYPES   = ["风格偏离", "重点不突出", "信息冗余", "低价值信息", "其他"]
SETTINGS_FILE       = str(ROOT / ".gui_settings.json")
MAX_MODELS          = 6

_prompt_cache: dict[str, str] = {}


# ── LLM 调用 ──────────────────────────────────────────────────────────

def load_prompt(prompt_file: str) -> str:
    if prompt_file in _prompt_cache:
        return _prompt_cache[prompt_file]
    p = Path(prompt_file)
    if not p.is_file():
        raise FileNotFoundError(f"Prompt 文件不存在: {prompt_file}")
    text = p.read_text(encoding="utf-8")
    _prompt_cache[prompt_file] = text
    return text



def call_llm(
    title: str,
    content: str,
    summary: str,
    prompt_file: str,
    api_key: str,
    model: str,
    api_base: str = "",
) -> dict:
    from openai import OpenAI

    system_prompt = load_prompt(prompt_file)
    user_msg = (
        f"【原文标题 | Original Title】: {title}\n"
        f"【原文内容 | Original Content】: {content}\n"
        f"【摘要 | Summary】: {summary}"
    )

    kwargs: dict = {"api_key": api_key}
    if api_base.strip():
        kwargs["base_url"] = api_base.strip()

    client = OpenAI(**kwargs)
    resp = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user",   "content": user_msg},
        ],
        temperature=0.0,
        response_format={"type": "json_object"},
    )
    return json.loads(resp.choices[0].message.content or "{}")


def parse_llm_result(raw: dict) -> ReviewerScore:
    try:
        score = int(raw.get("score", 0))
    except (ValueError, TypeError):
        score = 1
    score = max(1, min(5, score))

    raw_issues = raw.get("issue_types", [])
    if isinstance(raw_issues, str):
        raw_issues = [raw_issues]
    issues = [t for t in raw_issues if t in VALID_ISSUE_TYPES]

    comment = str(raw.get("comment", "") or "").strip()
    return ReviewerScore(score=score, issue_types=issues, comment=comment)


def eval_one(
    row: DataRow,
    model_entry: ModelEntry,
    cfg: dict,
    log_q: queue.Queue,
) -> tuple[DataRow, ModelEntry, ReviewerScore]:
    summary = (model_entry.summary or "").strip()
    if not summary:
        result = ReviewerScore(score=1, issue_types=["其他"], comment="模型无输出（空单元格）")
        log_q.put(f"[SKIP] id={row.id} {model_entry.model_col_name} => empty cell -> 1")
        return row, model_entry, result

    # NONE 及其他所有非空输出均交给 LLM judge（规则 8）

    try:
        raw = call_llm(
            title       = row.title,
            content     = row.content,
            summary     = summary,
            prompt_file = cfg["prompt_file"],
            api_key     = cfg["api_key"],
            model       = cfg["model"],
            api_base    = cfg.get("api_base", ""),
        )
        result = parse_llm_result(raw)
        log_q.put(
            f"[OK] id={row.id} {model_entry.model_col_name} "
            f"score={result.score} issues={result.issue_types}"
        )
    except Exception as e:
        result = ReviewerScore(score=1, issue_types=["其他"], comment=f"LLM调用失败: {e}")
        log_q.put(f"[ERR] id={row.id} {model_entry.model_col_name}: {e}")

    return row, model_entry, result


# ── 可折叠面板 ────────────────────────────────────────────────────────

class CollapsibleSection(ttk.Frame):
    """Drop-in collapsible panel — replaces ttk.LabelFrame.

    Children should be placed inside `self.content`.
    Pack the CollapsibleSection itself on the parent; set content_expand=True
    only for the section that should grow to fill available space.
    """

    def __init__(self, parent, text: str = "", expanded: bool = True,
                 content_expand: bool = False, **kwargs):
        super().__init__(parent, **kwargs)
        self._expanded       = expanded
        self._content_expand = content_expand

        # ── Header row ────────────────────────────────────────────────
        hdr = ttk.Frame(self)
        hdr.pack(fill="x", pady=(2, 0))

        self._arrow = ttk.Label(hdr, text="▼" if expanded else "▶",
                                width=2, cursor="hand2")
        self._arrow.pack(side="left", padx=(2, 0))

        title = ttk.Label(hdr, text=f" {text}",
                          font=("TkDefaultFont", 9, "bold"), cursor="hand2")
        title.pack(side="left")

        sep = ttk.Separator(hdr, orient="horizontal")
        sep.pack(side="left", fill="x", expand=True, padx=(6, 2), pady=5)

        # ── Content area ──────────────────────────────────────────────
        self.content = ttk.Frame(self, padding=(4, 2, 4, 4))
        if expanded:
            self.content.pack(fill="both", expand=content_expand)

        for w in (hdr, self._arrow, title, sep):
            w.bind("<Button-1>", self._toggle)

    def _toggle(self, _=None):
        if self._expanded:
            self.content.pack_forget()
            self._arrow.configure(text="▶")
        else:
            self.content.pack(fill="both", expand=self._content_expand)
            self._arrow.configure(text="▼")
        self._expanded = not self._expanded


# ── 单模型配置面板 ─────────────────────────────────────────────────────

class ModelConfigPanel(ttk.Frame):
    def __init__(self, parent, **kwargs):
        super().__init__(parent, **kwargs)
        self._build()

    def _build(self):
        ttk.Label(self, text="API Key:").grid(row=0, column=0, sticky="w", padx=6, pady=3)
        self.key_var = StringVar()
        ttk.Entry(self, textvariable=self.key_var, width=46, show="*").grid(
            row=0, column=1, columnspan=2, sticky="ew", padx=4)

        ttk.Label(self, text="模型:").grid(row=1, column=0, sticky="w", padx=6, pady=3)
        self.model_var = StringVar()
        self.model_combo = ttk.Combobox(self, textvariable=self.model_var, width=30)
        self.model_combo.grid(row=1, column=1, sticky="w", padx=4)
        self._fetch_btn = ttk.Button(self, text="检测模型", command=self._fetch_models)
        self._fetch_btn.grid(row=1, column=2, padx=4)

        ttk.Label(self, text="API Base URL:").grid(row=2, column=0, sticky="w", padx=6, pady=3)
        self.base_var = StringVar()
        ttk.Entry(self, textvariable=self.base_var, width=46).grid(
            row=2, column=1, columnspan=2, sticky="ew", padx=4)

        self.enabled_var = BooleanVar(value=True)
        ttk.Checkbutton(self, text="启用此模型", variable=self.enabled_var).grid(
            row=3, column=0, sticky="w", padx=6, pady=3)

        self.columnconfigure(1, weight=1)

    def _fetch_models(self):
        api_key  = self.key_var.get().strip()
        api_base = self.base_var.get().strip()
        if not api_key:
            messagebox.showerror("参数错误", "请先填写 API Key")
            return
        self._fetch_btn.configure(state="disabled", text="检测中...")

        def worker():
            try:
                from openai import OpenAI
                kw: dict = {"api_key": api_key}
                if api_base:
                    kw["base_url"] = api_base
                client = OpenAI(**kw)
                resp = client.models.list()
                model_ids = sorted(m.id for m in resp.data)
                self.after(0, lambda: self._on_fetch_done(model_ids, None))
            except Exception as e:
                self.after(0, lambda: self._on_fetch_done(None, str(e)))

        threading.Thread(target=worker, daemon=True).start()

    def _on_fetch_done(self, model_ids: list | None, error: str | None):
        self._fetch_btn.configure(state="normal", text="检测模型")
        if error:
            messagebox.showerror("检测失败", f"无法获取模型列表：\n{error}")
            return
        if model_ids:
            self.model_combo["values"] = model_ids
            if self.model_var.get() not in model_ids:
                self.model_var.set(model_ids[0])

    def get_cfg(self, prompt_file: str) -> dict | None:
        if not self.enabled_var.get():
            return None
        api_key = self.key_var.get().strip()
        model   = self.model_var.get().strip()
        if not api_key or not model:
            return None
        return {
            "api_key":     api_key,
            "model":       model,
            "api_base":    self.base_var.get().strip(),
            "prompt_file": prompt_file,
        }

    def to_dict(self) -> dict:
        return {
            "api_key": self.key_var.get(),
            "model":   self.model_var.get(),
            "api_base": self.base_var.get(),
            "enabled": self.enabled_var.get(),
            "models":  list(self.model_combo["values"]),
        }

    def from_dict(self, d: dict):
        self.key_var.set(d.get("api_key", ""))
        self.model_var.set(d.get("model", ""))
        self.base_var.set(d.get("api_base", ""))
        self.enabled_var.set(bool(d.get("enabled", True)))
        models = d.get("models", [])
        if models:
            self.model_combo["values"] = models


# ── 主界面 ────────────────────────────────────────────────────────────

class EvalApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("通知摘要 LLM 评测工具")
        self.resizable(True, True)
        self.minsize(840, 560)

        self._model_panels: list[ModelConfigPanel] = []
        self._log_texts:    list[scrolledtext.ScrolledText] = []
        self._model_pause_btns:  list[ttk.Button] = []
        self._model_stop_btns:   list[ttk.Button] = []
        self._model_status_lbls: list[ttk.Label]  = []
        self._log_q:        queue.Queue = queue.Queue()

        self._running     = False
        self._after_id    = None
        self._total_tasks = 0
        self._total_done  = 0
        self._done_lock   = threading.Lock()
        self._model_stop_events:  dict[int, threading.Event] = {}
        self._model_pause_events: dict[int, threading.Event] = {}

        # Tracks which prompt file is currently loaded in the editor
        self._editor_loaded_path: str = ""

        self._build_ui()
        self._load_settings()
        self.protocol("WM_DELETE_WINDOW", self._on_close)

    # ── 顶层 UI ───────────────────────────────────────────────────────

    def _build_ui(self):
        self._main_nb = ttk.Notebook(self)
        self._main_nb.pack(fill="both", expand=True, padx=4, pady=4)

        eval_tab    = ttk.Frame(self._main_nb)
        prompt_tab  = ttk.Frame(self._main_nb)
        report_tab  = ttk.Frame(self._main_nb)
        self._main_nb.add(eval_tab,   text="  评测  ")
        self._main_nb.add(prompt_tab, text="  Prompt 编辑  ")
        self._main_nb.add(report_tab, text="  结果报告  ")

        self._build_eval_tab(eval_tab)
        self._build_prompt_tab(prompt_tab)
        self._build_report_tab(report_tab)

        self._main_nb.bind("<<NotebookTabChanged>>", self._on_main_tab_changed)

    # ── 评测面板 ──────────────────────────────────────────────────────

    def _build_eval_tab(self, parent: ttk.Frame):
        # ── 预先创建日志区（不 pack），保证 _log_notebook 在 _add_model_tab 前存在 ──
        log_sec = CollapsibleSection(parent, text="运行日志",
                                     expanded=True, content_expand=True)
        self._log_notebook = ttk.Notebook(log_sec.content)
        all_frame = ttk.Frame(self._log_notebook)
        self._log_notebook.add(all_frame, text="全部")
        all_text = scrolledtext.ScrolledText(
            all_frame, height=8, state="disabled",
            font=("Consolas", 9), wrap="word")
        all_text.pack(fill="both", expand=True)
        self._log_texts = [all_text]

        # ── 输入文件 ──────────────────────────────────────────────────
        file_sec = CollapsibleSection(parent, text="输入文件 / 文件夹", expanded=True)
        file_sec.pack(fill="x", padx=6, pady=(4, 2))
        fc = file_sec.content   # alias for child widgets

        self._input_var = StringVar()
        ttk.Entry(fc, textvariable=self._input_var, width=56).grid(
            row=0, column=0, padx=4, pady=3, sticky="ew")
        ttk.Button(fc, text="选择文件",   command=self._pick_file,   width=10).grid(row=0, column=1, padx=3)
        ttk.Button(fc, text="选择文件夹", command=self._pick_folder, width=10).grid(row=0, column=2, padx=3)
        fc.columnconfigure(0, weight=1)

        ttk.Label(fc, text="Sheet:").grid(row=1, column=0, sticky="w", padx=4)
        self._sheet_var   = StringVar(value="自动检测")
        self._sheet_combo = ttk.Combobox(fc, textvariable=self._sheet_var,
                                         width=28, state="readonly")
        self._sheet_combo.grid(row=1, column=0, padx=56, sticky="w")
        ttk.Button(fc, text="刷新 Sheet", command=self._refresh_sheets, width=10).grid(
            row=1, column=1, padx=3)

        # ── 评测模型配置 ──────────────────────────────────────────────
        model_sec = CollapsibleSection(parent, text="评测模型配置", expanded=True)
        model_sec.pack(fill="x", padx=6, pady=2)
        mc = model_sec.content

        self._model_notebook = ttk.Notebook(mc)
        self._model_notebook.pack(fill="x", expand=False, padx=2, pady=(2, 0))

        ctrl_row = ttk.Frame(mc)
        ctrl_row.pack(fill="x", padx=2, pady=2)
        ttk.Button(ctrl_row, text="＋ 添加模型", command=self._add_model_tab,    width=12).pack(side="left", padx=3)
        ttk.Button(ctrl_row, text="－ 移除最后", command=self._remove_model_tab, width=12).pack(side="left", padx=2)

        prompt_row = ttk.Frame(mc)
        prompt_row.pack(fill="x", padx=2, pady=(0, 3))
        ttk.Label(prompt_row, text="Prompt 文件:").pack(side="left", padx=4)
        self._prompt_var = StringVar(value=PROMPT_FILE_DEFAULT)
        ttk.Entry(prompt_row, textvariable=self._prompt_var, width=46).pack(
            side="left", fill="x", expand=True, padx=3)
        ttk.Button(prompt_row, text="浏览", command=self._pick_prompt, width=6).pack(side="left", padx=3)

        for _ in range(3):
            self._add_model_tab(init=True)

        # ── 评测配置（默认折叠）────────────────────────────────────────
        cfg_sec = CollapsibleSection(parent, text="评测配置", expanded=False)
        cfg_sec.pack(fill="x", padx=6, pady=2)
        cc = cfg_sec.content

        ttk.Label(cc, text="每模型并发线程数:").grid(row=0, column=0, sticky="w", padx=4, pady=3)
        self._workers_var = IntVar(value=3)
        ttk.Spinbox(cc, from_=1, to=20, textvariable=self._workers_var, width=6).grid(
            row=0, column=1, sticky="w", padx=3)
        self._fmt_label = ttk.Label(cc, text="格式：检测中...", foreground="gray")
        self._fmt_label.grid(row=0, column=2, sticky="w", padx=16)

        ttk.Label(cc, text="输出目录:").grid(row=1, column=0, sticky="w", padx=4, pady=3)
        self._output_var = StringVar(value=OUTPUT_DIR_DEFAULT)
        ttk.Entry(cc, textvariable=self._output_var, width=38).grid(
            row=1, column=1, columnspan=2, sticky="ew", padx=3)
        ttk.Button(cc, text="浏览", command=self._pick_output, width=6).grid(row=1, column=3, padx=3)
        cc.columnconfigure(1, weight=1)

        # ── 进度条（常驻，不折叠）────────────────────────────────────
        prog_outer = ttk.Frame(parent)
        prog_outer.pack(fill="x", padx=6, pady=(3, 1))
        prog_row = ttk.Frame(prog_outer)
        prog_row.pack(fill="x")
        self._prog_var = IntVar(value=0)
        self._prog_bar = ttk.Progressbar(prog_row, variable=self._prog_var, maximum=100)
        self._prog_bar.pack(fill="x", side="left", expand=True)
        self._prog_label = ttk.Label(prog_row, text="0 / 0", width=12, anchor="e")
        self._prog_label.pack(side="right")
        self._status_label = ttk.Label(prog_outer, text="", foreground="gray")
        self._status_label.pack(anchor="w", padx=2)

        # ── 日志区（现在 pack，expand 填充剩余空间）──────────────────
        log_sec.pack(fill="both", expand=True, padx=6, pady=(2, 2))
        self._log_notebook.pack(fill="both", expand=True, padx=2, pady=2)

        # ── 按钮区（常驻）────────────────────────────────────────────
        btn_frame = ttk.Frame(parent)
        btn_frame.pack(fill="x", padx=6, pady=(2, 6))
        self._start_btn = ttk.Button(btn_frame, text="开始评测", command=self._start, width=14)
        self._start_btn.pack(side="left", padx=4)
        self._stop_btn = ttk.Button(btn_frame, text="全部停止", command=self._stop,
                                    width=10, state="disabled")
        self._stop_btn.pack(side="left", padx=4)
        self._pause_all_btn = ttk.Button(
            btn_frame, text="全部暂停", command=self._toggle_pause_all,
            width=10, state="disabled")
        self._pause_all_btn.pack(side="left", padx=4)
        ttk.Button(btn_frame, text="清空日志",    command=self._clear_log,    width=10).pack(side="left",  padx=4)
        ttk.Button(btn_frame, text="打开输出目录", command=self._open_output, width=14).pack(side="right", padx=4)

    # ── 结果报告面板 ──────────────────────────────────────────────────

    def _build_report_tab(self, parent: ttk.Frame):
        self._report_stats_a: list = []
        self._report_stats_list: list[tuple[str, list]] = []  # [(文件短名, stats), ...]

        # ── 文件选择区 ────────────────────────────────────────────────
        sel = ttk.LabelFrame(parent, text="数据文件")
        sel.pack(fill="x", padx=8, pady=(8, 4))

        # 参考版本（固定一行）
        ttk.Label(sel, text="参考版本:").grid(row=0, column=0, sticky="w", padx=6, pady=3)
        self._report_a_var = StringVar()
        ttk.Entry(sel, textvariable=self._report_a_var, width=52).grid(
            row=0, column=1, sticky="ew", padx=4)
        ttk.Button(sel, text="浏览", width=6,
                   command=lambda: self._pick_report_file(self._report_a_var)).grid(
            row=0, column=2, padx=4)

        ttk.Button(sel, text="生成报告", width=12, command=self._generate_report).grid(
            row=0, column=3, rowspan=99, padx=10, sticky="ns", pady=4)

        # 对比版本列表（动态行）
        self._cmp_rows_frame = ttk.Frame(sel)
        self._cmp_rows_frame.grid(row=1, column=0, columnspan=3, sticky="ew", padx=2, pady=(0, 2))
        self._cmp_vars: list[StringVar] = []

        btn_add = ttk.Button(sel, text="＋ 添加对比版本", width=16,
                             command=self._add_cmp_row)
        btn_add.grid(row=99, column=0, columnspan=3, sticky="w", padx=6, pady=(0, 4))

        sel.columnconfigure(1, weight=1)
        self._add_cmp_row()   # 默认给一个空行

        # ── 底部按钮（必须在 canvas 之前 pack，否则被 expand 挤掉）────────
        bot = ttk.Frame(parent)
        bot.pack(fill="x", padx=8, pady=(4, 8), side="bottom")
        ttk.Button(bot, text="导出报告 Excel", width=16,
                   command=self._export_report).pack(side="right", padx=4)

        # ── 报告主体（可滚动）────────────────────────────────────────
        vsb    = ttk.Scrollbar(parent, orient="vertical")
        canvas = tk.Canvas(parent, borderwidth=0, highlightthickness=0,
                           yscrollcommand=vsb.set)
        vsb.configure(command=canvas.yview)
        vsb.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True, padx=(8, 0))

        self._report_body = ttk.Frame(canvas)
        self._report_body_id = canvas.create_window((0, 0), window=self._report_body,
                                                     anchor="nw")

        def _sync_scroll(_e=None):
            canvas.configure(scrollregion=canvas.bbox("all"))

        def _sync_width(e):
            canvas.itemconfigure(self._report_body_id, width=e.width)

        self._report_body.bind("<Configure>", _sync_scroll)
        canvas.bind("<Configure>", _sync_width)

        def _on_wheel(e):
            canvas.yview_scroll(-1 * (e.delta // 120), "units")

        canvas.bind("<Enter>", lambda _e: canvas.bind_all("<MouseWheel>", _on_wheel))
        canvas.bind("<Leave>", lambda _e: canvas.unbind_all("<MouseWheel>"))

        self._report_hint = ttk.Label(
            self._report_body, text="← 选择文件后点击「生成报告」",
            foreground="gray")
        self._report_hint.pack(padx=20, pady=40)

    def _add_cmp_row(self):
        """在对比版本列表末尾追加一行。"""
        idx = len(self._cmp_vars)
        var = StringVar()
        self._cmp_vars.append(var)

        row_frame = ttk.Frame(self._cmp_rows_frame)
        row_frame.pack(fill="x", pady=1)

        ttk.Label(row_frame, text=f"对比版本 {idx + 1}:").pack(side="left", padx=(4, 2))
        ttk.Entry(row_frame, textvariable=var, width=48).pack(side="left", fill="x", expand=True, padx=2)
        ttk.Button(row_frame, text="浏览", width=6,
                   command=lambda v=var: self._pick_report_file(v)).pack(side="left", padx=2)
        ttk.Button(row_frame, text="×", width=3,
                   command=lambda f=row_frame, v=var: self._remove_cmp_row(f, v)).pack(side="left", padx=(0, 2))

    def _remove_cmp_row(self, frame: ttk.Frame, var: StringVar):
        """删除一行对比版本。"""
        if var in self._cmp_vars:
            self._cmp_vars.remove(var)
        frame.destroy()
        # 重新标号
        for i, child in enumerate(self._cmp_rows_frame.winfo_children()):
            lbl = child.winfo_children()[0]
            lbl.configure(text=f"对比版本 {i + 1}:")

    def _pick_report_file(self, var: StringVar):
        path = filedialog.askopenfilename(
            title="选择已评测的 Excel 文件",
            filetypes=[("Excel 文件", "*.xlsx *.xls"), ("所有文件", "*.*")])
        if path:
            var.set(path)

    def _generate_report(self):
        from core.report_stats import compute_stats

        path_a = self._report_a_var.get().strip()

        if not path_a or not Path(path_a).is_file():
            messagebox.showerror("参数错误", "请选择参考版本文件")
            return

        self._report_a_name = Path(path_a).name
        try:
            self._report_stats_a = compute_stats(path_a)
        except Exception as e:
            messagebox.showerror("读取失败", f"参考版本文件解析失败：\n{e}")
            return

        if not self._report_stats_a:
            from openpyxl import load_workbook as _lw
            try:
                _wb = _lw(path_a, read_only=False, data_only=True)
                _ws = _wb.worksheets[0]
                _hdrs = [str(c or "") for c in next(
                    _ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
                _wb.close()
                hint = "实际列头：" + ", ".join(_hdrs[:12]) if _hdrs else "第一行无内容"
            except Exception:
                hint = "无法读取列头"
            messagebox.showwarning(
                "提示",
                f"参考版本文件中未找到评分列（score_* 列），请确认选择的是评测输出文件。\n\n{hint}")
            return

        self._report_stats_list = []
        for var in self._cmp_vars:
            p = var.get().strip()
            if not p or not Path(p).is_file():
                continue
            short_name = Path(p).name
            try:
                stats = compute_stats(p)
                if stats:
                    self._report_stats_list.append((short_name, stats))
            except Exception as e:
                messagebox.showwarning("警告", f"对比版本文件解析失败（将跳过）：\n{p}\n{e}")

        self._render_report()

    def _render_report(self):
        for w in self._report_body.winfo_children():
            w.destroy()

        sa = self._report_stats_a
        a_name = getattr(self, "_report_a_name", "")
        cmp_list = self._report_stats_list   # [(file_name, stats), ...]

        model_names = [s.name for s in sa]

        if cmp_list:
            summary_sec = CollapsibleSection(
                self._report_body, text="综合对比", expanded=True)
            summary_sec.pack(fill="x", padx=4, pady=(4, 2))
            self._render_summary_table(summary_sec.content, sa, cmp_list, model_names)

        metrics_sec = CollapsibleSection(
            self._report_body, text="模型指标汇总", expanded=True)
        metrics_sec.pack(fill="x", padx=4, pady=(4, 2))
        self._render_metrics_table(metrics_sec.content, sa, a_name, cmp_list, model_names)

        attr_sec = CollapsibleSection(
            self._report_body, text="badcase 归因分析", expanded=True)
        attr_sec.pack(fill="x", padx=4, pady=(2, 4))
        self._render_attribution_table(attr_sec.content, sa, a_name, cmp_list, model_names)

    # ── 综合对比汇总表 ────────────────────────────────────────────────

    def _render_summary_table(self, parent, sa, cmp_list, model_names):
        a_map = {s.name: s for s in sa}

        cols      = ["对比版本", "模型数", "平均分 Δ均", "badcase 率 Δ均", "当前胜出", "综合结论"]
        col_ids   = [f"c{i}" for i in range(len(cols))]
        col_widths= [200, 60, 90, 110, 70, 120]
        anchors   = ["w", "center", "center", "center", "center", "center"]

        frame = ttk.Frame(parent)
        frame.pack(fill="x", padx=4, pady=4)

        tree = ttk.Treeview(frame, columns=col_ids, show="headings",
                            height=max(len(cmp_list), 1))
        for cid, cname, cw, anc in zip(col_ids, cols, col_widths, anchors):
            tree.heading(cid, text=cname)
            tree.column(cid, width=cw, anchor=anc, minwidth=50)

        tree.tag_configure("good",    foreground="#1a7a1a")
        tree.tag_configure("bad",     foreground="#cc3300")
        tree.tag_configure("neutral", foreground="#555555")

        for ver_name, sb in cmp_list:
            b_map  = {s.name: s for s in sb}
            shared = [n for n in model_names if n in b_map and n in a_map]
            n      = len(shared)

            if not shared:
                tree.insert("", END,
                            values=[ver_name, "0", "—", "—", "—", "无共同模型"],
                            tags=("neutral",))
                continue

            # Δ = 对比版本 - 参考版本（与明细表方向一致）
            avg_d  = sum(b_map[m].avg_score   - a_map[m].avg_score   for m in shared) / n
            rate_d = sum(b_map[m].badcase_rate - a_map[m].badcase_rate for m in shared) / n

            # 参考版本胜出（avg_score 更高）
            wins = sum(1 for m in shared if a_map[m].avg_score > b_map[m].avg_score)

            avg_str  = f"{'+' if avg_d  > 0 else ''}{avg_d:.2f}"
            rate_str = f"{'+' if rate_d > 0 else ''}{rate_d * 100:.1f}%"
            win_str  = f"{wins}/{n}"

            # 综合评分：avg_d<0 = 当前分更高(+1)；rate_d>0 = 当前 badcase 率更低(+1)
            points = 0
            if   avg_d  < -0.02: points += 1
            elif avg_d  >  0.02: points -= 1
            if   rate_d >  0.005: points += 1
            elif rate_d < -0.005: points -= 1

            conclusion, tag = {
                 2: ("参考版本更优",    "good"),
                 1: ("参考版本小幅占优", "good"),
                 0: ("两版本相当",      "neutral"),
                -1: ("参考版本小幅落后", "bad"),
                -2: ("参考版本较差",    "bad"),
            }[points]

            tree.insert("", END,
                        values=[ver_name, str(n), avg_str, rate_str, win_str, conclusion],
                        tags=(tag,))

        hsb = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(xscrollcommand=hsb.set)
        tree.pack(fill="x")
        hsb.pack(fill="x")

    # ── 指标汇总表 ────────────────────────────────────────────────────

    def _render_metrics_table(self, parent, sa, a_name, cmp_list, model_names):
        from core.report_stats import BADCASE_THRESHOLD

        # (label, ref值格式化, delta计算, higher_is_better)
        metrics = [
            ("模型平均分",
             lambda s: f"{s.avg_score:.2f}",
             lambda a, b: b.avg_score - a.avg_score,
             True),
            ("badcase 数",
             lambda s: str(s.badcase_count),
             lambda a, b: b.badcase_count - a.badcase_count,
             False),
            (f"badcase 率 (<{BADCASE_THRESHOLD}分)",
             lambda s: s.badcase_rate_pct,
             lambda a, b: (b.badcase_rate - a.badcase_rate) * 100,
             False),
        ]

        cur_col = a_name or "参考版本"
        # Δ 内嵌，对比版本不再单独列 Δ 列
        cols = ["指标", cur_col] + [ver_name for ver_name, _ in cmp_list]

        rows_data = []
        for model_name in model_names:
            sa_model = next((x for x in sa if x.name == model_name), None)
            if not sa_model:
                continue
            rows_data.append((
                [f"── {model_name} ──"] + [""] * (len(cols) - 1),
                None, None, True,
            ))
            for label, fmt, delta_fn, higher_is_better in metrics:
                row = [label, fmt(sa_model)]
                for _, sb in cmp_list:
                    b_map = {s.name: s for s in sb}
                    if model_name in b_map:
                        d    = delta_fn(sa_model, b_map[model_name])
                        sign = "+" if d > 0 else ""
                        val  = fmt(b_map[model_name])
                        if label.startswith("badcase 率"):
                            row.append(f"{val} ({sign}{d:.1f}%)")
                        elif isinstance(d, float):
                            row.append(f"{val} ({sign}{d:.2f})")
                        else:
                            row.append(f"{val} ({sign}{d})")
                    else:
                        row.append("")
                rows_data.append((row, label, higher_is_better, False))

        self._make_comparison_table(parent, cols, rows_data)

    # ── 归因分析表 ────────────────────────────────────────────────────

    def _render_attribution_table(self, parent, sa, a_name, cmp_list, model_names):
        from core.report_stats import VALID_ISSUE_TYPES

        attr_metrics = [("badcase 数", "badcase_count")] + \
                       [(t, t) for t in VALID_ISSUE_TYPES]

        cur_col = a_name or "参考版本"
        cols = ["归因类型", cur_col] + [ver_name for ver_name, _ in cmp_list]

        rows_data = []
        for model_name in model_names:
            sa_model = next((x for x in sa if x.name == model_name), None)
            if not sa_model:
                continue
            rows_data.append((
                [f"── {model_name} ──"] + [""] * (len(cols) - 1),
                None, None, True,
            ))
            for label, key in attr_metrics:
                val_a = (sa_model.badcase_count if key == "badcase_count"
                         else sa_model.issue_counts.get(key, 0))
                row = [label, str(val_a)]
                for _, sb in cmp_list:
                    b_map = {s.name: s for s in sb}
                    if model_name in b_map:
                        val_b = (b_map[model_name].badcase_count if key == "badcase_count"
                                 else b_map[model_name].issue_counts.get(key, 0))
                        d = val_b - val_a
                        sign = "+" if d > 0 else ""
                        row.append(f"{val_b} ({sign}{d})")
                    else:
                        row.append("")
                rows_data.append((row, label, False, False))

        self._make_comparison_table(parent, cols, rows_data)

    # ── 对比表构建助手 ────────────────────────────────────────────────

    def _make_comparison_table(self, parent, cols, rows_data):
        """
        版本横排对比表。
        rows_data: [(values, label, higher_is_better, is_header), ...]
          is_header=True  → 模型分组标题行（灰底加粗，无 Δ 着色）
          is_header=False → 数据行，higher_is_better 控制 Δ 颜色方向
        """
        frame = ttk.Frame(parent)
        frame.pack(fill="x", padx=4, pady=4)

        col_ids = [f"c{i}" for i in range(len(cols))]
        tree = ttk.Treeview(frame, columns=col_ids, show="headings",
                            height=len(rows_data))

        for i, (cid, cname) in enumerate(zip(col_ids, cols)):
            tree.heading(cid, text=cname)
            if i == 0:
                tree.column(cid, width=190, anchor="w",      minwidth=120)
            elif i == 1:
                tree.column(cid, width=110, anchor="center", minwidth=70)
            else:
                tree.column(cid, width=150, anchor="center", minwidth=100)

        tree.tag_configure("model_hdr", font=("TkDefaultFont", 9, "bold"),
                           background="#dde5f0")
        tree.tag_configure("pos_good", foreground="#1a7a1a")
        tree.tag_configure("neg_good", foreground="#1a7a1a")
        tree.tag_configure("pos_bad",  foreground="#cc3300")
        tree.tag_configure("neg_bad",  foreground="#cc3300")

        for values, _, higher_is_better, is_header in rows_data:
            if is_header:
                tree.insert("", END, values=values, tags=("model_hdr",))
            else:
                tag = "hib_true" if higher_is_better else "hib_false"
                tree.insert("", END, values=values, tags=(tag,))

        self._apply_delta_colors(tree, col_ids, cols, all_delta=False)

        hsb = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(xscrollcommand=hsb.set)
        tree.pack(fill="x")
        hsb.pack(fill="x")

    def _apply_delta_colors(self, tree, col_ids, cols, all_delta: bool = False):
        """对内嵌 Δ 的单元格着色，格式为 'value (±delta)'。"""
        import re
        if all_delta:
            delta_col_indices = list(range(1, len(cols)))
        else:
            # 跳过第 0 列（标签列）和第 1 列（参考版本列），其余对比版本列都需着色
            delta_col_indices = list(range(2, len(cols)))
        for item in tree.get_children():
            vals     = list(tree.item(item, "values"))
            cur_tags = set(tree.item(item, "tags"))
            if "model_hdr" in cur_tags:
                continue
            row_hib  = "hib_true" in cur_tags
            new_tags = cur_tags - {"hib_true", "hib_false"}
            for ci in delta_col_indices:
                if ci >= len(vals):
                    continue
                cell = str(vals[ci]).strip()
                # 从括号里提取 Δ 符号，如 "4.32 (+0.17)" 或 "12 (-2)"
                m = re.search(r'\(([+-])', cell)
                if not m:
                    continue
                is_positive = m.group(1) == "+"
                # 去掉括号内 Δ 为零的情况（如 (+0) / (+0.00) / (+0.0%)）
                zero_pat = re.search(r'\([+-]0(?:\.0+)?%?\)', cell)
                if zero_pat:
                    continue
                if is_positive:
                    new_tags.add("pos_good" if row_hib else "pos_bad")
                else:
                    new_tags.add("neg_good" if not row_hib else "neg_bad")
            tree.item(item, tags=tuple(new_tags))

    # ── 汇总一览 Sheet ────────────────────────────────────────────────

    def _write_summary_sheet(self, ws, sa, cmp_list, models, a_name,
                             HDR_FILL, GRP_FILL, GOOD_FILL, BAD_FILL,
                             BOLD, CENTER, LEFT, border):
        from core.report_stats import BADCASE_THRESHOLD
        from openpyxl.styles import PatternFill

        ref_label = a_name or "参考版本"

        # ── 构建列头 ──────────────────────────────────────────────────
        # 参考版本三列 + 每个对比版本三列
        base_cols  = ["平均分", "badcase 数", f"badcase 率(<{BADCASE_THRESHOLD}分)"]
        cmp_names  = [ver_name for ver_name, _ in cmp_list]

        header_row1 = ["模型名", ref_label] + [""] * (len(base_cols) - 1)
        header_row2 = ["",       "平均分", "badcase 数", f"badcase 率(<{BADCASE_THRESHOLD}分)"]
        for ver_name in cmp_names:
            header_row1 += [ver_name] + [""] * (len(base_cols) - 1)
            header_row2 += ["平均分", "badcase 数", f"badcase 率(<{BADCASE_THRESHOLD}分)"]

        # 写第一行（版本分组标题，跨列合并）
        for ci, v in enumerate(header_row1, 1):
            cell = ws.cell(1, ci, v if v else "")
            cell.font      = BOLD
            cell.fill      = HDR_FILL
            cell.alignment = CENTER
            cell.border    = border
        # 合并参考版本列头
        ref_start = 2
        ws.merge_cells(start_row=1, start_column=ref_start,
                       end_row=1, end_column=ref_start + len(base_cols) - 1)
        for idx, ver_name in enumerate(cmp_names):
            cs = ref_start + len(base_cols) + idx * len(base_cols)
            ws.merge_cells(start_row=1, start_column=cs,
                           end_row=1, end_column=cs + len(base_cols) - 1)

        # 写第二行（指标子列名）
        for ci, v in enumerate(header_row2, 1):
            cell = ws.cell(2, ci, v)
            cell.font      = BOLD
            cell.fill      = HDR_FILL
            cell.alignment = CENTER
            cell.border    = border

        # 设列宽
        from openpyxl.utils import get_column_letter
        ws.column_dimensions["A"].width = 28
        col_letter_idx = 2
        for _ in range(1 + len(cmp_names)):
            for w in (12, 12, 16):
                ws.column_dimensions[get_column_letter(col_letter_idx)].width = w
                col_letter_idx += 1

        # ── 写数据行 ──────────────────────────────────────────────────
        b_maps = {v: {s.name: s for s in sb} for v, sb in cmp_list}
        data_row = 3

        for model_name in models:
            ref = next((x for x in sa if x.name == model_name), None)
            if not ref:
                continue

            c = 1
            # 模型名
            cell = ws.cell(data_row, c, model_name); c += 1
            cell.font = BOLD; cell.fill = GRP_FILL
            cell.alignment = LEFT; cell.border = border

            # 参考版本三列（纯数值，无 Δ）
            ws.cell(data_row, c, ref.avg_score).number_format    = "0.00";  ws.cell(data_row, c).alignment = CENTER; ws.cell(data_row, c).border = border; c += 1
            ws.cell(data_row, c, ref.badcase_count).alignment    = CENTER;  ws.cell(data_row, c).border = border; c += 1
            ws.cell(data_row, c, ref.badcase_rate).number_format = "0.0%";  ws.cell(data_row, c).alignment = CENTER; ws.cell(data_row, c).border = border; c += 1

            # 每个对比版本三列（数值 + Δ 内嵌文本）
            for ver_name in cmp_names:
                bm  = b_maps.get(ver_name, {})
                cmp = bm.get(model_name)

                if cmp is None:
                    for _ in range(len(base_cols)):
                        cell = ws.cell(data_row, c, "—")
                        cell.alignment = CENTER; cell.border = border; c += 1
                    continue

                # 平均分
                d_avg  = round(cmp.avg_score - ref.avg_score, 2)
                sign   = "+" if d_avg > 0 else ""
                cell   = ws.cell(data_row, c,
                                 f"{cmp.avg_score:.2f} ({sign}{d_avg:.2f})")
                cell.alignment = CENTER; cell.border = border
                cell.fill = (GOOD_FILL if d_avg > 0.005
                             else BAD_FILL if d_avg < -0.005
                             else PatternFill())
                c += 1

                # badcase 数
                d_bc  = cmp.badcase_count - ref.badcase_count
                sign  = "+" if d_bc > 0 else ""
                cell  = ws.cell(data_row, c,
                                f"{cmp.badcase_count} ({sign}{d_bc})")
                cell.alignment = CENTER; cell.border = border
                cell.fill = (BAD_FILL if d_bc > 0
                             else GOOD_FILL if d_bc < 0
                             else PatternFill())
                c += 1

                # badcase 率
                d_rate = round(cmp.badcase_rate - ref.badcase_rate, 4)
                sign   = "+" if d_rate > 0 else ""
                cell   = ws.cell(data_row, c,
                                 f"{cmp.badcase_rate*100:.1f}% ({sign}{d_rate*100:.1f}%)")
                cell.alignment = CENTER; cell.border = border
                cell.fill = (BAD_FILL if d_rate > 0.005
                             else GOOD_FILL if d_rate < -0.005
                             else PatternFill())
                c += 1

            data_row += 1

        # 冻结前两行
        ws.freeze_panes = "A3"

    # ── 导出报告 Excel ────────────────────────────────────────────────

    def _export_report(self):
        if not self._report_stats_a:
            messagebox.showinfo("提示", "请先生成报告")
            return

        path = filedialog.asksaveasfilename(
            title="保存报告",
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")],
            initialfile=f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        if not path:
            return

        try:
            self._write_report_excel(path)
            messagebox.showinfo("完成", f"报告已导出：\n{path}")
        except Exception as e:
            messagebox.showerror("导出失败", str(e))

    def _write_report_excel(self, path: str):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from core.report_stats import VALID_ISSUE_TYPES, BADCASE_THRESHOLD

        wb       = Workbook()
        sa       = self._report_stats_a
        cmp_list = self._report_stats_list
        a_name   = getattr(self, "_report_a_name", "参考版本")
        models   = [s.name for s in sa]

        # ── 样式常量 ──────────────────────────────────────────────────────
        HDR_FILL  = PatternFill("solid", fgColor="D9E1F2")   # 列标题
        GRP_FILL  = PatternFill("solid", fgColor="DDE5F0")   # 模型分组行
        GOOD_FILL = PatternFill("solid", fgColor="C6EFCE")   # 好的 Δ（绿）
        BAD_FILL  = PatternFill("solid", fgColor="FFC7CE")   # 坏的 Δ（红）
        BOLD      = Font(bold=True)
        CENTER    = Alignment(horizontal="center", vertical="center", wrap_text=True)
        LEFT      = Alignment(horizontal="left",   vertical="center")
        thin      = Side(style="thin")
        border    = Border(left=thin, right=thin, top=thin, bottom=thin)

        def _hdr(ws, r, c, v):
            cell = ws.cell(r, c, v)
            cell.font, cell.fill, cell.alignment, cell.border = BOLD, HDR_FILL, CENTER, border

        def _grp(ws, r, c, v):
            cell = ws.cell(r, c, v)
            cell.font, cell.fill, cell.alignment, cell.border = BOLD, GRP_FILL, LEFT,   border

        def _dat(ws, r, c, v, align=CENTER, fmt=None):
            cell = ws.cell(r, c, v)
            cell.alignment, cell.border = align, border
            if fmt:
                cell.number_format = fmt

        # ── 列结构：[参考版本, ver1, ver2, …]（Δ 内嵌到对比版本单元格，不单独列）──
        cur_col  = a_name or "参考版本"
        # col_info[i] = (col_name, ver_name)  ver_name="" 表示参考版本列
        col_info: list[tuple[str, str]] = [(cur_col, "")]
        for ver_name, _ in cmp_list:
            col_info.append((ver_name, ver_name))

        b_maps = {v: {s.name: s for s in sb} for v, sb in cmp_list}

        # ── 通用写 Sheet 函数 ─────────────────────────────────────────────
        def write_sheet(ws, label_col, rows, get_val, get_delta,
                        lower_better_fn, rate_fn, first_w=22):
            """
            rows            : [(label, key), ...]
            get_val(stats, key)  -> 数值
            get_delta(va, vb)    -> 差值（数字）
            lower_better_fn(key) -> bool
            rate_fn(key)         -> bool（True 表示百分比列）
            """
            from openpyxl.utils import get_column_letter
            # 第 1 行：列标题
            _hdr(ws, 1, 1, label_col)
            ws.column_dimensions["A"].width = first_w
            for ci, (cname, _) in enumerate(col_info, 2):
                _hdr(ws, 1, ci, cname)
                ws.column_dimensions[get_column_letter(ci)].width = 20

            row = 2
            for model_name in models:
                sa_model = next((x for x in sa if x.name == model_name), None)
                if not sa_model:
                    continue

                # 模型分组标题行
                _grp(ws, row, 1, f"── {model_name} ──")
                for ci in range(2, len(col_info) + 2):
                    ws.cell(row, ci).fill   = GRP_FILL
                    ws.cell(row, ci).border = border
                row += 1

                for label, key in rows:
                    _dat(ws, row, 1, label, LEFT)
                    val_a = get_val(sa_model, key)
                    r_fmt = "0.0%" if rate_fn(key) else None
                    _dat(ws, row, 2, val_a, fmt=r_fmt)

                    for ci, (_, ver_name) in enumerate(col_info[1:], 3):
                        bm = b_maps.get(ver_name, {})
                        if model_name not in bm:
                            _dat(ws, row, ci, "—")
                            continue
                        val_b = get_val(bm[model_name], key)
                        d     = get_delta(val_a, val_b)
                        sign  = "+" if d > 0 else ""
                        if rate_fn(key):
                            text = f"{val_b * 100:.1f}% ({sign}{d * 100:.1f}%)"
                        elif isinstance(d, float):
                            text = f"{val_b:.2f} ({sign}{d:.2f})"
                        else:
                            text = f"{val_b} ({sign}{d})"
                        cell = ws.cell(row, ci, text)
                        cell.alignment, cell.border = CENTER, border
                        lower = lower_better_fn(key)
                        is_good = (d < 0) if lower else (d > 0)
                        threshold = 0.005 if isinstance(d, float) else 0
                        if abs(d) > threshold:
                            cell.fill = GOOD_FILL if is_good else BAD_FILL
                    row += 1

        # ── Sheet 0：汇总一览 ─────────────────────────────────────────────
        ws0 = wb.active
        ws0.title = "汇总一览"
        self._write_summary_sheet(ws0, sa, cmp_list, models, a_name,
                                  HDR_FILL, GRP_FILL, GOOD_FILL, BAD_FILL,
                                  BOLD, CENTER, LEFT, border)

        # ── Sheet 1：模型指标汇总 ─────────────────────────────────────────
        ws1 = wb.create_sheet("模型指标汇总")

        metric_rows = [
            ("模型平均分",                           "avg_score"),
            ("badcase 数",                           "badcase_count"),
            (f"badcase 率 (<{BADCASE_THRESHOLD}分)", "badcase_rate"),
        ]
        write_sheet(
            ws1, "指标", metric_rows,
            get_val        = lambda s, k: getattr(s, k),
            get_delta      = lambda va, vb: round(vb - va, 4),
            lower_better_fn= lambda k: k != "avg_score",
            rate_fn        = lambda k: k == "badcase_rate",
            first_w        = 26,
        )

        # ── Sheet 2：badcase 归因分析 ─────────────────────────────────────
        ws2 = wb.create_sheet("badcase 归因分析")

        attr_rows = [("badcase 数", "badcase_count")] + [(t, t) for t in VALID_ISSUE_TYPES]
        write_sheet(
            ws2, "归因类型", attr_rows,
            get_val        = lambda s, k: (s.badcase_count if k == "badcase_count"
                                           else s.issue_counts.get(k, 0)),
            get_delta      = lambda va, vb: vb - va,
            lower_better_fn= lambda k: True,
            rate_fn        = lambda k: False,
            first_w        = 18,
        )

        wb.save(path)

    # ── Prompt 编辑面板 ───────────────────────────────────────────────

    def _build_prompt_tab(self, parent: ttk.Frame):
        # 顶栏：路径 + 操作按钮
        header = ttk.Frame(parent)
        header.pack(fill="x", padx=8, pady=(8, 2))

        ttk.Label(header, text="文件路径:").pack(side="left", padx=(0, 4))
        self._editor_path_lbl = ttk.Label(
            header, text="", foreground="#336699",
            relief="sunken", padding=(4, 1), anchor="w")
        self._editor_path_lbl.pack(side="left", fill="x", expand=True, padx=(0, 6))

        ttk.Button(header, text="重新加载", command=self._reload_prompt_editor, width=10).pack(side="right", padx=2)
        ttk.Button(header, text="保存",     command=self._save_prompt_editor,   width=8).pack(side="right", padx=2)

        # 字数统计栏
        stats_bar = ttk.Frame(parent)
        stats_bar.pack(fill="x", padx=8, pady=(0, 2))
        self._editor_stats_lbl = ttk.Label(stats_bar, text="", foreground="gray")
        self._editor_stats_lbl.pack(side="right")
        self._editor_modified_lbl = ttk.Label(stats_bar, text="", foreground="#cc6600")
        self._editor_modified_lbl.pack(side="left")

        # 编辑器主体
        self._prompt_editor = scrolledtext.ScrolledText(
            parent, font=("Consolas", 10), wrap="word", undo=True)
        self._prompt_editor.pack(fill="both", expand=True, padx=8, pady=(0, 8))
        self._prompt_editor.bind("<<Modified>>", self._on_editor_modified)

    def _on_main_tab_changed(self, _event=None):
        """When user switches to the Prompt editor tab, sync content."""
        selected = self._main_nb.index(self._main_nb.select())
        if selected == 1:
            current_path = self._prompt_var.get().strip()
            if current_path != self._editor_loaded_path:
                self._reload_prompt_editor()

    def _reload_prompt_editor(self):
        path = self._prompt_var.get().strip()
        if not path or not Path(path).is_file():
            self._editor_path_lbl.configure(text=path or "(未设置)")
            self._prompt_editor.configure(state="normal")
            self._prompt_editor.delete("1.0", END)
            self._prompt_editor.configure(state="disabled")
            self._editor_stats_lbl.configure(text="文件不存在")
            return
        try:
            text = Path(path).read_text(encoding="utf-8")
            self._prompt_editor.configure(state="normal")
            self._prompt_editor.delete("1.0", END)
            self._prompt_editor.insert("1.0", text)
            self._prompt_editor.edit_modified(False)
            self._editor_loaded_path = path
            self._editor_path_lbl.configure(text=path)
            self._editor_modified_lbl.configure(text="")
            self._update_editor_stats()
        except Exception as e:
            messagebox.showerror("读取失败", f"无法读取文件：\n{e}")

    def _save_prompt_editor(self):
        path = self._editor_loaded_path
        if not path:
            path = self._prompt_var.get().strip()
        if not path:
            messagebox.showerror("保存失败", "没有关联的文件路径")
            return
        try:
            content = self._prompt_editor.get("1.0", END)
            Path(path).write_text(content, encoding="utf-8")
            # 清除 prompt 缓存，使下次评测重新读取
            _prompt_cache.pop(path, None)
            self._prompt_editor.edit_modified(False)
            self._editor_modified_lbl.configure(text="")
            self._editor_loaded_path = path
            self._editor_path_lbl.configure(text=path)
            self._update_editor_stats()
        except Exception as e:
            messagebox.showerror("保存失败", f"写入文件失败：\n{e}")

    def _on_editor_modified(self, _event=None):
        if self._prompt_editor.edit_modified():
            self._editor_modified_lbl.configure(text="● 未保存")
            self._update_editor_stats()

    def _update_editor_stats(self):
        try:
            content = self._prompt_editor.get("1.0", END)
            words = len(content.split())
            chars = len(content.replace(" ", "").replace("\n", "").replace("\t", ""))
            lines = content.count("\n")
            self._editor_stats_lbl.configure(
                text=f"{words} 词  |  {chars} 字符  |  {lines} 行")
        except Exception:
            pass

    # ── 模型/日志 tab 管理 ────────────────────────────────────────────

    def _add_model_tab(self, init: bool = False):
        n = len(self._model_panels) + 1
        if n > MAX_MODELS:
            if not init:
                messagebox.showinfo("提示", f"最多支持 {MAX_MODELS} 个模型")
            return

        panel = ModelConfigPanel(self._model_notebook)
        self._model_notebook.add(panel, text=f"模型 {n}")
        self._model_panels.append(panel)

        log_frame = ttk.Frame(self._log_notebook)
        self._log_notebook.add(log_frame, text=f"模型 {n}")

        # ── 每模型控制栏 ──────────────────────────────────────────────
        model_idx_ref = n - 1   # 0-based index for this model

        ctrl = ttk.Frame(log_frame)
        ctrl.pack(fill="x", padx=2, pady=(2, 0))

        status_lbl = ttk.Label(ctrl, text="待机", foreground="gray", width=8)
        status_lbl.pack(side="left", padx=(4, 8))

        pause_btn = ttk.Button(
            ctrl, text="暂停", width=6, state="disabled",
            command=lambda i=model_idx_ref: self._toggle_pause_model(i))
        pause_btn.pack(side="left", padx=2)

        stop_btn = ttk.Button(
            ctrl, text="停止", width=6, state="disabled",
            command=lambda i=model_idx_ref: self._stop_model(i))
        stop_btn.pack(side="left", padx=2)

        ttk.Separator(log_frame, orient="horizontal").pack(fill="x", pady=(2, 0))

        log_text = scrolledtext.ScrolledText(
            log_frame, height=8, state="disabled",
            font=("Consolas", 9), wrap="word")
        log_text.pack(fill="both", expand=True)

        self._log_texts.append(log_text)
        self._model_pause_btns.append(pause_btn)
        self._model_stop_btns.append(stop_btn)
        self._model_status_lbls.append(status_lbl)

    def _remove_model_tab(self):
        if len(self._model_panels) <= 1:
            messagebox.showinfo("提示", "至少保留一个模型")
            return
        self._model_notebook.forget(self._model_notebook.tabs()[-1])
        self._log_notebook.forget(self._log_notebook.tabs()[-1])
        self._model_panels.pop()
        self._log_texts.pop()
        self._model_pause_btns.pop()
        self._model_stop_btns.pop()
        self._model_status_lbls.pop()

    # ── 文件选择 ──────────────────────────────────────────────────────

    def _pick_file(self):
        path = filedialog.askopenfilename(
            title="选择输入文件",
            filetypes=[("Excel/CSV 文件", "*.xlsx *.xls *.csv"), ("所有文件", "*.*")])
        if path:
            self._input_var.set(path)
            self._refresh_sheets()

    def _pick_folder(self):
        path = filedialog.askdirectory(title="选择包含评测文件的文件夹")
        if path:
            self._input_var.set(path)
            self._sheet_combo["values"] = ["自动检测"]
            self._sheet_var.set("自动检测")

    def _pick_prompt(self):
        path = filedialog.askopenfilename(
            title="选择 Prompt 文件",
            filetypes=[("Markdown/文本", "*.md *.txt"), ("所有文件", "*.*")])
        if path:
            self._prompt_var.set(path)

    def _pick_output(self):
        path = filedialog.askdirectory(title="选择输出目录")
        if path:
            self._output_var.set(path)

    def _refresh_sheets(self):
        path = self._input_var.get().strip()
        if not path or not Path(path).is_file():
            return
        ext = Path(path).suffix.lower()
        if ext not in {".xlsx", ".xls"}:
            self._sheet_combo["values"] = ["自动检测"]
            self._sheet_var.set("自动检测")
            return
        try:
            sheets = detect_sheets(path)
            self._sheet_combo["values"] = ["自动检测"] + sheets
            auto = detect_data_sheet(path)
            self._sheet_var.set(auto)
            self._detect_and_show_format(path, auto)
        except Exception as e:
            self._log_all(f"[WARN] 读取 sheet 列表失败: {e}")

    def _detect_and_show_format(self, file_path: str, sheet_name: str):
        try:
            from openpyxl import load_workbook as _lw
            wb = _lw(file_path, read_only=True, data_only=True)
            ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]
            headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            wb.close()
            if _is_compact_format(headers):
                wb2 = _lw(file_path, read_only=True, data_only=True)
                ws2 = wb2[sheet_name] if sheet_name in wb2.sheetnames else wb2.worksheets[0]
                n_models = ws2.max_column - 5
                wb2.close()
                self._fmt_label.configure(
                    text=f"格式A（紧凑）：{n_models} 个模型列，输出时自动插入评分列",
                    foreground="#1a7a1a")
            else:
                self._fmt_label.configure(
                    text="格式B（含评分列）：直接写入 score/标签/comment",
                    foreground="#1a4a9a")
        except Exception:
            self._fmt_label.configure(text="格式：未知", foreground="gray")

    # ── 日志工具 ──────────────────────────────────────────────────────

    def _append_to_text(self, widget: scrolledtext.ScrolledText, msg: str):
        ts = datetime.now().strftime("%H:%M:%S")
        line = f"[{ts}] {msg}\n"
        widget.configure(state="normal")
        widget.insert(END, line)
        widget.see(END)
        widget.configure(state="disabled")

    def _log_all(self, msg: str):
        self._append_to_text(self._log_texts[0], msg)

    def _log(self, msg: str):
        self._log_all(msg)

    def _clear_log(self):
        for t in self._log_texts:
            t.configure(state="normal")
            t.delete("1.0", END)
            t.configure(state="disabled")

    def _poll_log(self):
        try:
            while True:
                item = self._log_q.get_nowait()
                if isinstance(item, tuple):
                    model_idx, msg = item
                    self._append_to_text(self._log_texts[0], f"[M{model_idx + 1}] {msg}")
                    tab_idx = model_idx + 1
                    if tab_idx < len(self._log_texts):
                        self._append_to_text(self._log_texts[tab_idx], msg)
                else:
                    self._append_to_text(self._log_texts[0], item)
        except queue.Empty:
            pass
        if self._running:
            self._after_id = self.after(200, self._poll_log)

    # ── 进度 ──────────────────────────────────────────────────────────

    def _set_status(self, text: str):
        self._status_label.configure(text=text)

    def _set_progress(self, done: int, total: int):
        pct = int(done / total * 100) if total else 0
        self._prog_var.set(pct)
        self._prog_label.configure(text=f"{done} / {total}")

    def _inc_progress(self):
        with self._done_lock:
            self._total_done += 1
            done, total = self._total_done, self._total_tasks
        self.after(0, lambda: self._set_progress(done, total))

    # ── 开始 / 停止 ───────────────────────────────────────────────────

    def _start(self):
        if self._running:
            return

        input_path  = self._input_var.get().strip()
        prompt_file = self._prompt_var.get().strip()
        output_dir  = self._output_var.get().strip()
        workers     = self._workers_var.get()
        sheet_sel   = self._sheet_var.get().strip()

        if not input_path:
            messagebox.showerror("参数错误", "请选择输入文件或文件夹")
            return
        if not Path(prompt_file).is_file():
            messagebox.showerror("参数错误", f"Prompt 文件不存在:\n{prompt_file}")
            return

        model_cfgs: list[tuple[int, dict]] = []
        for i, panel in enumerate(self._model_panels):
            cfg = panel.get_cfg(prompt_file)
            if cfg:
                model_cfgs.append((i, cfg))

        if not model_cfgs:
            messagebox.showerror("参数错误", "请至少启用并配置一个模型（需填写 API Key 和模型名）")
            return

        files = self._collect_files(input_path)
        if not files:
            messagebox.showerror("参数错误", "未找到可处理的 Excel/CSV 文件")
            return

        self._total_done  = 0
        self._total_tasks = 0
        self._running     = True
        self._start_btn.configure(state="disabled")
        self._stop_btn.configure(state="normal")
        self._pause_all_btn.configure(state="normal", text="全部暂停")

        # 初始化每模型的独立停止/暂停事件，并启用对应控制按钮
        self._model_stop_events.clear()
        self._model_pause_events.clear()
        for panel_idx, _ in model_cfgs:
            self._model_stop_events[panel_idx]  = threading.Event()
            pe = threading.Event()
            pe.set()   # SET = 运行中（未暂停）
            self._model_pause_events[panel_idx] = pe
            if panel_idx < len(self._model_pause_btns):
                self._model_pause_btns[panel_idx].configure(state="normal", text="暂停")
                self._model_stop_btns[panel_idx].configure(state="normal")
                self._model_status_lbls[panel_idx].configure(
                    text="运行中", foreground="#1a7a1a")

        enabled_names = [f"模型{i+1}({c['model']})" for i, c in model_cfgs]
        self._log_all(
            f"开始评测，共 {len(files)} 个文件，"
            f"启用模型: {', '.join(enabled_names)}，每模型并发: {workers}"
        )
        self._poll_log()

        threading.Thread(
            target=self._run_all,
            args=(files, sheet_sel, model_cfgs, workers, output_dir),
            daemon=True,
        ).start()

    def _stop(self):
        """全部停止 — 设置所有模型的停止事件。"""
        if self._running:
            for se in self._model_stop_events.values():
                se.set()
            for pe in self._model_pause_events.values():
                pe.set()   # 恢复暂停中的模型，使其能检测到停止信号
            self._log_all("[用户] 已发送全部停止信号，等待当前任务完成...")

    def _toggle_pause_all(self):
        """全部暂停 / 全部继续切换。"""
        # 判断当前是否所有模型都已暂停
        all_paused = all(
            not pe.is_set()
            for pe in self._model_pause_events.values()
        )
        if all_paused:
            # 全部继续
            for i, pe in self._model_pause_events.items():
                pe.set()
                if i < len(self._model_pause_btns):
                    se = self._model_stop_events.get(i, threading.Event())
                    if not se.is_set():
                        self._model_pause_btns[i].configure(text="暂停")
                        self._model_status_lbls[i].configure(
                            text="运行中", foreground="#1a7a1a")
            self._pause_all_btn.configure(text="全部暂停")
            self._log_all("[用户] 全部继续")
        else:
            # 全部暂停（跳过已停止的模型）
            for i, pe in self._model_pause_events.items():
                se = self._model_stop_events.get(i, threading.Event())
                if not se.is_set():
                    pe.clear()
                    if i < len(self._model_pause_btns):
                        self._model_pause_btns[i].configure(text="继续")
                        self._model_status_lbls[i].configure(
                            text="已暂停", foreground="#cc6600")
            self._pause_all_btn.configure(text="全部继续")
            self._log_all("[用户] 全部暂停")


        """停止单个模型。"""
        se = self._model_stop_events.get(model_idx)
        pe = self._model_pause_events.get(model_idx)
        if se:
            se.set()
        if pe:
            pe.set()   # 若已暂停，先恢复以让线程检测到停止信号
        if model_idx < len(self._model_status_lbls):
            self._model_status_lbls[model_idx].configure(
                text="停止中", foreground="#cc3300")
            self._model_pause_btns[model_idx].configure(state="disabled")
            self._model_stop_btns[model_idx].configure(state="disabled")
        self._log_q.put((model_idx, f"[用户] 已发送停止信号"))

    def _toggle_pause_model(self, model_idx: int):
        """暂停 / 继续 单个模型。"""
        pe = self._model_pause_events.get(model_idx)
        if not pe:
            return
        btn = self._model_pause_btns[model_idx] if model_idx < len(self._model_pause_btns) else None
        lbl = self._model_status_lbls[model_idx] if model_idx < len(self._model_status_lbls) else None
        if pe.is_set():          # 当前运行中 → 暂停
            pe.clear()
            if btn: btn.configure(text="继续")
            if lbl: lbl.configure(text="已暂停", foreground="#cc6600")
            self._log_q.put((model_idx, "[用户] 已暂停，等待当前 LLM 调用完成后暂停提交"))
        else:                    # 当前暂停 → 继续
            pe.set()
            if btn: btn.configure(text="暂停")
            if lbl: lbl.configure(text="运行中", foreground="#1a7a1a")
            self._log_q.put((model_idx, "[用户] 已恢复运行"))

    def _model_finished(self, model_idx: int):
        """某一模型线程结束后，在主线程中更新其按钮状态（GUI 安全）。"""
        if model_idx < len(self._model_pause_btns):
            self._model_pause_btns[model_idx].configure(state="disabled", text="暂停")
            self._model_stop_btns[model_idx].configure(state="disabled")
            stopped = self._model_stop_events.get(model_idx, threading.Event()).is_set()
            lbl_text = "已停止" if stopped else "完成"
            lbl_color = "#cc3300" if stopped else "gray"
            self._model_status_lbls[model_idx].configure(
                text=lbl_text, foreground=lbl_color)

    def _on_finish(self, success: bool, msg: str):
        self._running = False
        self._start_btn.configure(state="normal")
        self._stop_btn.configure(state="disabled")
        self._pause_all_btn.configure(state="disabled", text="全部暂停")
        for btn in self._model_pause_btns:
            btn.configure(state="disabled", text="暂停")
        for btn in self._model_stop_btns:
            btn.configure(state="disabled")
        self._set_status("")
        if self._after_id:
            self.after_cancel(self._after_id)
        try:
            while True:
                item = self._log_q.get_nowait()
                if isinstance(item, tuple):
                    mi, text = item
                    self._append_to_text(self._log_texts[0], f"[M{mi+1}] {text}")
                    if mi + 1 < len(self._log_texts):
                        self._append_to_text(self._log_texts[mi + 1], text)
                else:
                    self._log_all(item)
        except queue.Empty:
            pass
        if success:
            messagebox.showinfo("完成", msg)
        else:
            messagebox.showerror("出错", msg)

    # ── 文件收集 ──────────────────────────────────────────────────────

    def _collect_files(self, input_path: str) -> list[Path]:
        p = Path(input_path)
        if p.is_file():
            return [p] if p.suffix.lower() in VALID_EXTENSIONS else []
        if p.is_dir():
            result = []
            for ext in VALID_EXTENSIONS:
                result.extend(p.rglob(f"*{ext}"))
            return sorted(result)
        return []

    # ── 核心评测流程 ──────────────────────────────────────────────────

    def _run_all(
        self,
        files      : list[Path],
        sheet_sel  : str,
        model_cfgs : list[tuple[int, dict]],
        workers    : int,
        output_dir : str,
    ):
        try:
            self.after(0, lambda: self._set_status("读取文件中..."))

            file_data:  dict[Path, tuple] = {}
            base_tasks: list[tuple] = []

            for fpath in files:
                if all(se.is_set() for se in self._model_stop_events.values()):
                    break
                try:
                    sheet_name = (
                        detect_data_sheet(str(fpath))
                        if sheet_sel in ("", "自动检测")
                        else sheet_sel
                    )
                    reader = ExcelReader(str(fpath), sheet_name)
                    rows, warnings = reader.load()
                    for w in warnings:
                        self._log_q.put(f"[WARN] {fpath.name}: {w}")

                    fmt_str = "格式A（紧凑）" if reader.is_compact else "格式B（含评分列）"
                    self._log_q.put(f"  {fpath.name}: {fmt_str}，{len(rows)} 行")

                    file_data[fpath] = (rows, sheet_name, reader.is_compact)
                    for row in rows:
                        for me in row.models:
                            base_tasks.append((fpath, row, me))
                except Exception as e:
                    self._log_q.put(f"[ERR] 读取 {fpath.name} 失败: {e}")

            total_tasks = len(base_tasks) * len(model_cfgs)
            self._total_tasks = total_tasks
            self._log_q.put(
                f"共 {len(base_tasks)} 条数据 × {len(model_cfgs)} 个模型 "
                f"= {total_tasks} 个评测任务"
            )
            self.after(0, lambda: self._set_progress(0, total_tasks))
            self.after(0, lambda: self._set_status("评测中..."))

            os.makedirs(output_dir, exist_ok=True)

            threads = [
                threading.Thread(
                    target=self._run_model,
                    args=(model_idx, cfg, file_data, base_tasks, workers, output_dir),
                    daemon=True,
                )
                for model_idx, cfg in model_cfgs
            ]
            for t in threads:
                t.start()
            for t in threads:
                t.join()

            msg = f"评测完成！共处理 {self._total_done}/{total_tasks} 条，{len(model_cfgs)} 个模型。"
            self.after(0, lambda: self._on_finish(True, msg))

        except Exception as e:
            self.after(0, lambda: self._on_finish(False, f"评测过程出错:\n{e}"))

    def _run_model(
        self,
        model_idx  : int,
        cfg        : dict,
        file_data  : dict[Path, tuple],
        base_tasks : list[tuple],
        workers    : int,
        output_dir : str,
    ):
        stop_event  = self._model_stop_events[model_idx]
        pause_event = self._model_pause_events[model_idx]  # SET=运行, CLEAR=暂停

        def mlog(msg: str):
            self._log_q.put((model_idx, msg))

        model_name = cfg["model"]
        mlog(f"=== 开始评测，judge 模型: {model_name} ===")

        handlers: dict[Path, tuple] = {}
        for fpath, (_, sheet_name, is_compact) in file_data.items():
            handler = (ExcelOutputBuilder(str(fpath), sheet_name)
                       if is_compact
                       else ExcelWriter(str(fpath), sheet_name))
            handlers[fpath] = (is_compact, handler)

        local_q: queue.Queue = queue.Queue()

        def drain_local():
            try:
                while True:
                    self._log_q.put((model_idx, local_q.get_nowait()))
            except queue.Empty:
                pass

        remaining = list(base_tasks)
        pending: dict = {}     # future → (fpath, row, me)

        def fill():
            """提交新任务直到池满（暂停或停止时不提交）。"""
            while (not stop_event.is_set()
                   and pause_event.is_set()
                   and len(pending) < workers
                   and remaining):
                fpath, row, me = remaining.pop(0)
                f = executor.submit(eval_one, row, me, cfg, local_q)
                pending[f] = (fpath, row, me)

        with ThreadPoolExecutor(max_workers=workers) as executor:
            fill()

            while pending or remaining:
                pause_event.wait()       # 暂停时在此阻塞，恢复后继续

                if stop_event.is_set():
                    executor.shutdown(wait=False, cancel_futures=True)
                    pending.clear()
                    remaining.clear()
                    mlog("[已停止]")
                    break

                fill()                   # 暂停恢复后补充任务

                if not pending:
                    continue             # remaining 未空但均未提交（暂停刚恢复时）

                done, _ = cf_wait(pending, return_when=FIRST_COMPLETED, timeout=0.5)

                for future in done:
                    fpath, _, _ = pending.pop(future)
                    drain_local()
                    try:
                        row_obj, me_result, score_result = future.result()
                        is_compact, handler = handlers[fpath]
                        if is_compact:
                            handler.add_result(
                                row_index = row_obj.row_index,
                                src_col   = me_result.src_col,
                                result    = score_result,
                            )
                        else:
                            handler.write_scores(
                                row_index    = row_obj.row_index,
                                model_col    = me_result.src_col,
                                reviewer_idx = 0,
                                score        = score_result.score or 1,
                                issue_types  = score_result.issue_types,
                                comment      = score_result.comment or "",
                            )
                    except Exception as e:
                        mlog(f"[ERR] 写入结果失败: {e}")

                    self._inc_progress()

                fill()   # 处理完一批后立即补充

        drain_local()

        ts         = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_model = re.sub(r'[\\/:*?"<>|]', '_', model_name)
        for fpath, (is_compact, handler) in handlers.items():
            out_name = f"{fpath.stem}_eval_{safe_model}_{ts}.xlsx"
            out_path = str(Path(output_dir) / out_name)
            mlog(f"[SAVE] 正在保存 {out_name} ...")
            try:
                handler.save(out_path)
                mlog(f"[SAVE] ✓ 已保存: {out_name}")
            except Exception as e:
                mlog(f"[ERR] 保存失败: {e}")

        mlog(f"=== 模型 {model_name} 全部完成 ===")
        self.after(0, lambda: self._model_finished(model_idx))

    # ── 打开输出目录 ──────────────────────────────────────────────────

    def _open_output(self):
        path = self._output_var.get().strip()
        if path and Path(path).is_dir():
            os.startfile(path)
        else:
            messagebox.showinfo("提示", f"目录不存在: {path}")

    # ── 设置持久化 ────────────────────────────────────────────────────

    def _load_settings(self):
        try:
            with open(SETTINGS_FILE, encoding="utf-8") as f:
                s = json.load(f)
        except Exception:
            return
        if s.get("input"):   self._input_var.set(s["input"])
        if s.get("output"):  self._output_var.set(s["output"])
        if s.get("prompt"):  self._prompt_var.set(s["prompt"])
        if s.get("workers"): self._workers_var.set(s["workers"])
        if s.get("report_a"): self._report_a_var.set(s["report_a"])
        for i, p in enumerate(s.get("report_cmp", [])):
            if not p:
                continue
            while i >= len(self._cmp_vars):
                self._add_cmp_row()
            self._cmp_vars[i].set(p)

        for i, d in enumerate(s.get("model_configs", [])):
            if i < len(self._model_panels):
                self._model_panels[i].from_dict(d)
            else:
                self._add_model_tab()
                self._model_panels[-1].from_dict(d)

    def _save_settings(self):
        s = {
            "input":        self._input_var.get(),
            "output":       self._output_var.get(),
            "prompt":       self._prompt_var.get(),
            "workers":      self._workers_var.get(),
            "report_a":     self._report_a_var.get(),
            "report_cmp":   [v.get() for v in self._cmp_vars],
            "model_configs": [p.to_dict() for p in self._model_panels],
        }
        try:
            with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
                json.dump(s, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def _on_close(self):
        self._save_settings()
        self.destroy()


# ── 入口 ──────────────────────────────────────────────────────────────

def main():
    app = EvalApp()
    app.mainloop()


if __name__ == "__main__":
    main()
