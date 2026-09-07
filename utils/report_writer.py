"""
report_writer.py
----------------
Excel 报告输出模块：将 EvalReport 写入格式化的 Excel 文件。

输出结构（多 Sheet）：
  Sheet 1 - 摘要 (Summary)   : 整体统计 + 各规则通过率
  Sheet 2 - 明细 (Detail)    : 每条记录 × 每条规则的评测结果
  Sheet 3 - 失败记录 (Failed): 仅包含存在 error 的记录
"""

from __future__ import annotations

import os
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import (
    Alignment, Font, PatternFill, Border, Side
)
from openpyxl.utils import get_column_letter

from core.evaluator import EvalReport


# ------------------------------------------------------------------
# 颜色常量
# ------------------------------------------------------------------
COLOR_HEADER   = "2F5496"   # 深蓝（表头背景）
COLOR_PASS     = "C6EFCE"   # 浅绿（通过）
COLOR_FAIL     = "FFC7CE"   # 浅红（失败）
COLOR_WARNING  = "FFEB9C"   # 浅黄（警告）
COLOR_SUBHEAD  = "D9E1F2"   # 浅蓝（子标题）


def _make_fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _make_border() -> Border:
    thin = Side(style="thin", color="BFBFBF")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _style_header_row(ws, row: int, col_count: int) -> None:
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.font      = Font(bold=True, color="FFFFFF", size=11)
        cell.fill      = _make_fill(COLOR_HEADER)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = _make_border()


def _auto_col_width(ws) -> None:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                val_len = len(str(cell.value)) if cell.value else 0
                max_len = max(max_len, val_len)
            except Exception:  # noqa: BLE001
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 60)


# ------------------------------------------------------------------
# 主类
# ------------------------------------------------------------------

class ReportWriter:
    def __init__(self, output_path: str):
        """
        Parameters
        ----------
        output_path : 输出 Excel 文件路径（含文件名）
        """
        self.output_path = output_path

    def write(self, report: EvalReport) -> str:
        """将报告写入 Excel 并返回实际输出路径。
        若 output_path 是目录，则自动生成带时间戳的文件名。
        """
        path = self._resolve_path()
        os.makedirs(os.path.dirname(path), exist_ok=True)

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            self._write_summary(writer, report)
            self._write_detail(writer, report)
            self._write_failed(writer, report)

        # 二次打开做样式美化
        self._apply_styles(path, report)
        return path

    # ------------------------------------------------------------------
    # 路径处理
    # ------------------------------------------------------------------
    def _resolve_path(self) -> str:
        if os.path.isdir(self.output_path) or not self.output_path.endswith(".xlsx"):
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            return os.path.join(self.output_path, f"eval_report_{ts}.xlsx")
        return self.output_path

    # ------------------------------------------------------------------
    # Sheet 1: 摘要
    # ------------------------------------------------------------------
    def _write_summary(self, writer: pd.ExcelWriter, report: EvalReport) -> None:
        s = report.summary

        # 整体统计
        overview = pd.DataFrame([
            {"指标": "总记录数",     "值": s.total_records},
            {"指标": "通过记录数",   "值": s.passed_records},
            {"指标": "失败记录数",   "值": s.failed_records},
            {"指标": "通过率 (%)",   "值": s.pass_rate},
            {"指标": "总错误数",     "值": s.total_errors},
            {"指标": "总警告数",     "值": s.total_warnings},
            {"指标": "评测耗时 (s)", "值": s.elapsed_seconds},
        ])

        # 规则统计
        rule_rows = []
        for rid, stat in s.rule_stats.items():
            total = stat["pass"] + stat["fail"]
            rule_rows.append({
                "规则 ID"  : rid,
                "规则名称" : stat["name"],
                "通过数"   : stat["pass"],
                "失败数"   : stat["fail"],
                "通过率(%)": round(stat["pass"] / total * 100, 2) if total else 0,
            })
        rule_df = pd.DataFrame(rule_rows)

        overview.to_excel(writer, sheet_name="摘要", index=False, startrow=1)
        if not rule_df.empty:
            rule_df.to_excel(
                writer, sheet_name="摘要", index=False,
                startrow=len(overview) + 4
            )

    # ------------------------------------------------------------------
    # Sheet 2: 明细
    # ------------------------------------------------------------------
    def _write_detail(self, writer: pd.ExcelWriter, report: EvalReport) -> None:
        rows = []
        for rr in report.records:
            for res in rr.rule_results:
                row = {
                    "行号"    : rr.row_index,
                    "规则 ID" : res["rule_id"],
                    "规则名称": res["rule_name"],
                    "字段"    : res["field"],
                    "严重级别": res["severity"],
                    "是否通过": "✓ 通过" if res["passed"] else "✗ 失败",
                    "说明"    : res["message"],
                }
                # 附加原始字段值（方便排查问题）
                field = res["field"]
                row["字段值"] = rr.record.get(field, "") if field else ""
                rows.append(row)

        df = pd.DataFrame(rows)
        df.to_excel(writer, sheet_name="明细", index=False)

    # ------------------------------------------------------------------
    # Sheet 3: 失败记录
    # ------------------------------------------------------------------
    def _write_failed(self, writer: pd.ExcelWriter, report: EvalReport) -> None:
        rows = []
        for rr in report.records:
            if not rr.passed:
                failed_rules = [
                    f"{r['rule_name']}: {r['message']}"
                    for r in rr.rule_results
                    if not r["passed"] and r["severity"] == "error"
                ]
                row = {k: v for k, v in rr.record.items() if not k.startswith("_")}
                row["错误数"]   = rr.error_count
                row["警告数"]   = rr.warning_count
                row["失败规则"] = " | ".join(failed_rules)
                rows.append(row)

        df = pd.DataFrame(rows) if rows else pd.DataFrame(columns=["（无失败记录）"])
        df.to_excel(writer, sheet_name="失败记录", index=False)

    # ------------------------------------------------------------------
    # 样式美化
    # ------------------------------------------------------------------
    def _apply_styles(self, path: str, report: EvalReport) -> None:
        wb = load_workbook(path)

        # --- 摘要 Sheet ---
        ws = wb["摘要"]
        ws.insert_rows(1)
        ws["A1"] = "📊 评测摘要报告"
        ws["A1"].font      = Font(bold=True, size=14, color="2F5496")
        ws["A1"].alignment = Alignment(horizontal="left")
        _style_header_row(ws, 2, 2)
        _auto_col_width(ws)

        # --- 明细 Sheet ---
        ws2 = wb["明细"]
        _style_header_row(ws2, 1, ws2.max_column)
        for row in ws2.iter_rows(min_row=2):
            passed_cell = row[5]  # "是否通过" 列（索引 5）
            fill = _make_fill(COLOR_PASS) if "通过" in str(passed_cell.value) else _make_fill(COLOR_FAIL)
            for cell in row:
                cell.fill   = fill
                cell.border = _make_border()
                cell.alignment = Alignment(wrap_text=True, vertical="center")
        _auto_col_width(ws2)

        # --- 失败记录 Sheet ---
        ws3 = wb["失败记录"]
        _style_header_row(ws3, 1, ws3.max_column)
        for row in ws3.iter_rows(min_row=2):
            for cell in row:
                cell.fill   = _make_fill(COLOR_FAIL)
                cell.border = _make_border()
                cell.alignment = Alignment(wrap_text=True, vertical="center")
        _auto_col_width(ws3)

        wb.save(path)
